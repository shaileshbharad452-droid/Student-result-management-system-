import openpyxl
import os

FILE_NAME = "student_results.xlsx"


# Create Excel file
def create_excel_file():
    if not os.path.exists(FILE_NAME):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Student Results"

        headers = [
            "Roll No", "Name", "Class",
            "Subject 1", "Subject 2", "Subject 3",
            "Subject 4", "Subject 5",
            "Total", "Percentage", "Grade", "Status"
        ]

        sheet.append(headers)
        workbook.save(FILE_NAME)


# Calculate result
def calculate_result(marks):
    total = sum(marks)
    percentage = total / 5

    if percentage >= 90:
        grade = "A+"
    elif percentage >= 80:
        grade = "A"
    elif percentage >= 70:
        grade = "B"
    elif percentage >= 60:
        grade = "C"
    elif percentage >= 50:
        grade = "D"
    else:
        grade = "F"

    if all(mark >= 35 for mark in marks):
        status = "PASS"
    else:
        status = "FAIL"
        grade = "F"

    return total, percentage, grade, status


# Add student
def add_student():
    print("\n----- ADD STUDENT RESULT -----")

    roll_no = input("Enter Roll No: ")
    name = input("Enter Student Name: ")
    course = input("Enter Class/Course: ")

    marks = []

    for i in range(1, 6):
        mark = float(input(f"Enter marks of Subject {i}: "))
        marks.append(mark)

    total, percentage, grade, status = calculate_result(marks)

    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    row = [
        roll_no, name, course,
        marks[0], marks[1], marks[2],
        marks[3], marks[4],
        total, percentage, grade, status
    ]

    sheet.append(row)
    workbook.save(FILE_NAME)

    print("\nStudent result saved successfully!")
    print("Total:", total)
    print("Percentage:", f"{percentage:.2f}%")
    print("Grade:", grade)
    print("Status:", status)


# Get student result
def get_result():
    print("\n----- GET STUDENT RESULT -----")

    roll_no = input("Enter Roll No: ")

    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if str(row[0]) == roll_no:

            print("\n--------------------------------")
            print("        STUDENT RESULT")
            print("--------------------------------")
            print("Roll No     :", row[0])
            print("Name        :", row[1])
            print("Class       :", row[2])
            print("Total       :", row[8])
            print("Percentage  :", f"{row[9]:.2f}%")
            print("Grade       :", row[10])
            print("Status      :", row[11])
            print("--------------------------------")

            found = True
            break

    if not found:
        print("Student not found!")


# Show all student data
def show_all_data():
    print("\n----- ALL STUDENT DATA -----")

    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    if sheet.max_row <= 1:
        print("No student records found.")
        return

    print("-" * 85)

    print(
        f"{'Roll':<8}"
        f"{'Name':<15}"
        f"{'Class':<10}"
        f"{'Total':<8}"
        f"{'Percent':<12}"
        f"{'Grade':<8}"
        f"{'Status':<8}"
    )

    print("-" * 85)

    for row in sheet.iter_rows(min_row=2, values_only=True):
        print(
            f"{str(row[0]):<8}"
            f"{str(row[1]):<15}"
            f"{str(row[2]):<10}"
            f"{str(row[8]):<8}"
            f"{row[9]:<12.2f}"
            f"{str(row[10]):<8}"
            f"{str(row[11]):<8}"
        )

    print("-" * 85)


# Menu
def menu():
    create_excel_file()

    while True:
        print("\n========================================")
        print("      STUDENT RESULT MANAGEMENT")
        print("========================================")
        print("1. Add Student Result")
        print("2. Get Student Result")
        print("3. Show All Student Data")
        print("4. Exit")
        print("========================================")

        choice = input("Enter your choice: ")

        if choice == "1":
            add_student()

        elif choice == "2":
            get_result()

        elif choice == "3":
            show_all_data()

        elif choice == "4":
            print("Thank you! Program closed.")
            break

        else:
            print("Invalid choice! Please try again.")


# Start program
menu()